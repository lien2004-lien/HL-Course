
import re
import unicodedata


from io import BytesIO




from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, response
from django.shortcuts import get_object_or_404
from django.core.exceptions import PermissionDenied

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from apps.courses.models import CourseClass
from apps.payments.models import Payment


@login_required
def export_class_students(request, class_id):

    if request.user.role != "teacher":
        raise PermissionDenied(
            "Chỉ giảng viên mới được xuất danh sách."
        )

    class_group = get_object_or_404(
        CourseClass.objects.select_related("course"),
        id=class_id,
        teacher=request.user,
    )

    enrollments = (
        class_group.enrollments
        .filter(payment__status=Payment.Status.PAID)
        .select_related("student")
        .order_by("student__username")
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Danh sach hoc vien"

    worksheet.merge_cells("A1:I1")
    worksheet["A1"] = "HL COURSE - DANH SÁCH HỌC VIÊN"
    worksheet["A1"].font = Font(
        size=16,
        bold=True,
        color="FFFFFF",
    )
    worksheet["A1"].fill = PatternFill(
        "solid",
        fgColor="1E40AF",
    )
    worksheet["A1"].alignment = Alignment(
        horizontal="center"
    )
    worksheet.row_dimensions[1].height = 32

    worksheet["A3"] = "Khóa học:"
    worksheet["B3"] = class_group.course.title

    worksheet["A4"] = "Lớp:"
    worksheet["B4"] = class_group.name

    worksheet["A5"] = "Giảng viên:"
    worksheet["B5"] = request.user.get_full_name() or request.user.username

    worksheet["A6"] = "Sĩ số:"
    worksheet["B6"] = len(enrollments)

    headers = [
        "STT",
        "Mã học viên",
        "Tên đăng nhập",
        "Họ tên",
        "Email",
        "Khóa học",
        "Lớp học",
        "Ngày đăng ký",
        "Thanh toán",
    ]

    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(
            row=8,
            column=column,
            value=header,
        )
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.fill = PatternFill(
            "solid",
            fgColor="2563EB",
        )
        cell.alignment = Alignment(
            horizontal="center",
        )

    for index, enrollment in enumerate(enrollments, start=1):

        student = enrollment.student
        row = index + 8

        values = [
            index,
            student.id,
            student.username,
            student.get_full_name() or "Chưa cập nhật",
            student.email,
            class_group.course.title,
            class_group.name,
            enrollment.enrolled_at.replace(tzinfo=None),
            "Đã thanh toán",
        ]

        for column, value in enumerate(values, start=1):
            worksheet.cell(
                row=row,
                column=column,
                value=value,
            )

        worksheet.cell(row=row, column=8).number_format = (
            "dd/mm/yyyy hh:mm"
        )

    widths = [8, 14, 22, 28, 32, 30, 18, 22, 20]

    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[
            get_column_letter(index)
        ].width = width

    worksheet.freeze_panes = "A9"
    worksheet.auto_filter.ref = (
        f"A8:I{max(9, worksheet.max_row)}"
    )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )


    def safe_filename(value):
        value = unicodedata.normalize("NFKD", value)
        value = "".join(
            char for char in value
            if not unicodedata.combining(char)
        )
    
        value = value.replace("đ", "d").replace("Đ", "D")
        value = re.sub(r"[^A-Za-z0-9]+", "_", value)
        return value.strip("_")


    course_name = safe_filename(class_group.course.title)
    class_name = safe_filename(class_group.name)

    filename = f"HL_Course_{course_name}_{class_name}.xlsx"

    response["Content-Disposition"] = (
    f'attachment; filename="{filename}"'
)

    return response
